"""Deterministic structural validation for review workbook templates."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from create_review_templates import CONTROLLED_VOCABULARY, REQUIRED_SHEETS


REQUIRED_HEADERS = {
    "PRISMA筛选记录": ["阶段", "数量", "排除原因", "备注"],
    "文献批判性编码": ["ID", "题名", "暴露/可达性指标", "空间尺度", "正义维度"],
    "Claim-Evidence Map": [
        "核心主张",
        "证据来源",
        "Claim_ID",
        "Evidence_IDs",
        "Argument_IDs",
        "Evidence_State",
    ],
    "Evidence Trace": [
        "Evidence_ID",
        "Paper_ID",
        "Citation_Key_or_DOI",
        "Full_Text_Status",
        "Claim_or_Finding",
        "Evidence_State",
        "Source_Locator",
        "Source_Section",
        "Source_Page",
        "Source_Table",
        "Source_Figure",
        "Source_Supplement",
        "Support_Type",
        "Reviewer_Status",
    ],
    "跨论文论证地图": [
        "Argument_ID",
        "Review_Question",
        "Synthesis_Claim",
        "Supporting_Evidence_IDs",
        "Contradictory_Evidence_IDs",
        "Evidence_Consistency",
        "Evidence_Strength",
        "Unresolved_Gap",
    ],
    "段落Synthesis Contract": [
        "Paragraph_ID",
        "Section",
        "Paragraph_Function",
        "Target_Synthesis_Claim",
        "Argument_IDs",
        "Evidence_IDs",
        "Forbidden_Escalation",
        "Citation_Audit_Status",
        "Ready_to_Draft",
    ],
    "期刊定位": ["目标期刊", "适配主题", "必须强调", "避免"],
    "图表规划": ["图表编号", "要证明的主张", "数据来源", "图表类型"],
}


def _headers(ws):
    return [cell.value for cell in ws[1]]


def _validate_unique_headers(ws):
    headers = _headers(ws)
    nonempty = [header for header in headers if header not in (None, "")]
    duplicates = sorted({header for header in nonempty if nonempty.count(header) > 1})
    if duplicates:
        raise ValueError(f"{ws.title}: duplicate headers: {duplicates}")


def _validate_vocab_bindings(ws):
    headers = set(_headers(ws))
    validations = list(ws.data_validations.dataValidation)
    formulas = {validation.formula1 for validation in validations}
    for field, values in CONTROLLED_VOCABULARY.items():
        if field not in headers:
            continue
        expected = '"' + ",".join(values) + '"'
        if expected not in formulas:
            raise ValueError(f"{ws.title}: missing controlled validation for {field}")


def _validate_contract_rows(workbook):
    evidence = workbook["Evidence Trace"]
    evidence_headers = _headers(evidence)
    evidence_columns = {header: index + 1 for index, header in enumerate(evidence_headers)}
    full_text_status = evidence_columns["Full_Text_Status"]
    locator_columns = [
        evidence_columns[field]
        for field in (
            "Source_Section",
            "Source_Page",
            "Source_Table",
            "Source_Figure",
            "Source_Supplement",
        )
    ]
    for row in range(2, evidence.max_row + 1):
        if evidence.cell(row=row, column=full_text_status).value != "ABSTRACT_ONLY":
            continue
        if any(evidence.cell(row=row, column=column).value not in (None, "") for column in locator_columns):
            raise ValueError(
                f"Evidence Trace row {row}: ABSTRACT_ONLY cannot contain full-text locator fields"
            )

    paragraphs = workbook["段落Synthesis Contract"]
    paragraph_headers = _headers(paragraphs)
    paragraph_columns = {header: index + 1 for index, header in enumerate(paragraph_headers)}
    target_claim = paragraph_columns["Target_Synthesis_Claim"]
    evidence_ids = paragraph_columns["Evidence_IDs"]
    ready = paragraph_columns["Ready_to_Draft"]
    for row in range(2, paragraphs.max_row + 1):
        claim = paragraphs.cell(row=row, column=target_claim).value
        linked_evidence = paragraphs.cell(row=row, column=evidence_ids).value
        ready_value = paragraphs.cell(row=row, column=ready).value
        if claim not in (None, "") and linked_evidence in (None, "") and ready_value == "YES":
            raise ValueError(
                f"段落Synthesis Contract row {row}: factual/synthesis claim requires Evidence_IDs"
            )


def validate_workbook(path: Path) -> None:
    if not path.is_file():
        raise FileNotFoundError(path)

    workbook = load_workbook(path, read_only=False, data_only=False)
    reopened = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet_names = set(workbook.sheetnames)
        missing_sheets = sorted(REQUIRED_SHEETS - sheet_names)
        if missing_sheets:
            raise ValueError(f"missing required sheets: {missing_sheets}")

        for sheet_name, required in REQUIRED_HEADERS.items():
            ws = workbook[sheet_name]
            _validate_unique_headers(ws)
            headers = set(_headers(ws))
            missing_headers = [header for header in required if header not in headers]
            if missing_headers:
                raise ValueError(f"{sheet_name}: missing headers: {missing_headers}")
            _validate_vocab_bindings(ws)

        _validate_contract_rows(workbook)

        if set(reopened.sheetnames) != sheet_names:
            raise ValueError("workbook sheet names changed on reopen")
    finally:
        workbook.close()
        reopened.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    validate_workbook(args.workbook.resolve())
    print(f"VALIDATION_PASS {args.workbook.resolve()}")


if __name__ == "__main__":
    main()
